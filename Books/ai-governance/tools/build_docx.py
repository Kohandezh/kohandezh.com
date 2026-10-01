#!/usr/bin/env python3
"""Phase 20 — editable Persian DOCX (MasterPrompt §22, §45, §75, §118). python-docx, native RTL, A4, real Word styles, TOC field, custom provenance metadata.
Output: docx/fa/book-fa.docx"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Pt, Mm, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import db, render  # noqa: E402

PUB = yaml.safe_load(open(ROOT / "config" / "publishing.yaml", encoding="utf-8"))
PROV = yaml.safe_load(open(ROOT / "config" / "provenance.yaml", encoding="utf-8"))
SRC = {d["doc_key"]: d for d in json.load(open(ROOT / "manifests" / "sources.json", encoding="utf-8"))["documents"]}
FA_FONT, LATIN_FONT, MONO_FONT = "Vazirmatn", "Inter", "Noto Sans Mono"
BOX_STYLES = {"editorial_synthesis": ("KDJ Editorial Note", "FFF8E6", "E0A800"), "editorial_localization_ir": ("KDJ Localization", "EAF7EE", "2E8B57"),
              "executive": ("KDJ Executive Box", "EEF2FF", "4054B2"), "technical": ("KDJ Technical Box", "F3EEFC", "7B3FBF"), "template": ("KDJ Template", "FDF0F0", "C0392B"),
              "checklist": ("KDJ Checklist", "F6F8FA", "5B6470"), "scenario": ("KDJ Iranian Scenario", "FFF3E0", "E67E22")}

# --------------------------------------------------------------------------- low-level RTL helpers
def set_rtl(p, rtl=True):
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:bidi",):
        el = pPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); pPr.append(el)
        el.set(qn("w:val"), "1" if rtl else "0")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    return p

def set_run_fonts(run, fa=FA_FONT, latin=LATIN_FONT, size=None, bold=None, mono=False, color=None):
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    f_lat = MONO_FONT if mono else latin; f_cs = MONO_FONT if mono else fa
    rf.set(qn("w:ascii"), f_lat); rf.set(qn("w:hAnsi"), f_lat); rf.set(qn("w:cs"), f_cs); rf.set(qn("w:eastAsia"), f_lat)
    if not mono:
        rtl = rPr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl"); rPr.append(rtl)
        rtl.set(qn("w:val"), "1")
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rPr.append(lang)
    lang.set(qn("w:bidi"), "fa-IR"); lang.set(qn("w:val"), "en-US")
    if size: run.font.size = Pt(size); rPr_sz = rPr.find(qn("w:szCs")) or OxmlElement("w:szCs"); rPr_sz.set(qn("w:val"), str(int(size * 2))); rPr.append(rPr_sz)
    if bold is not None: run.font.bold = bold; b = OxmlElement("w:bCs"); b.set(qn("w:val"), "1" if bold else "0"); rPr.append(b)
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(p, fill: str, border: str):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); pPr.append(shd)
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "18" if side == "right" else "4"); el.set(qn("w:space"), "4"); el.set(qn("w:color"), border); bdr.append(el)
    pPr.append(bdr)

def add_inline(p, text: str, size=None, base_bold=False):
    """**bold** and `code` inline markup → runs."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); set_run_fonts(r, size=size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1]); set_run_fonts(r, size=size or 9.5, mono=True)
        else:
            r = p.add_run(tok); set_run_fonts(r, size=size, bold=True if base_bold else None)

def para(doc, text="", style=None, size=None, bold=False, align=None, fill=None, border=None, mono=False, space_after=6):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    set_rtl(p)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if mono:
        r = p.add_run(text); set_run_fonts(r, size=size or 9, mono=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pPr = p._p.get_or_add_pPr(); bidi = pPr.find(qn("w:bidi")); bidi.set(qn("w:val"), "0")
    elif text:
        add_inline(p, text, size=size, base_bold=bold)
    if fill: shade(p, fill, border or "999999")
    return p

def heading(doc, text: str, level: int, cid: str | None = None):
    h = doc.add_heading(level=level); set_rtl(h)
    r = h.add_run(text); set_run_fonts(r, size={0: 26, 1: 20, 2: 16, 3: 13.5, 4: 12}.get(level, 12), bold=True, color="1B1F24")
    if cid:
        para(doc, cid, size=8, mono=True, space_after=2)
    return h

def add_table(doc, header: list[str], rows: list[list[str]], fill=None):
    n = max(len(header), max((len(r) for r in rows), default=0))
    t = doc.add_table(rows=1, cols=n); t.style = "Table Grid"
    tblPr = t._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); tblPr.append(bv)
    for i in range(n):
        c = t.rows[0].cells[i]; c.text = ""; p = c.paragraphs[0]; set_rtl(p); add_inline(p, header[i] if i < len(header) else "", size=9.5, base_bold=True)
        tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill or "E8ECF0"); tcPr.append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i in range(n):
            cells[i].text = ""; p = cells[i].paragraphs[0]; set_rtl(p); add_inline(p, row[i] if i < len(row) else "", size=9.5)
    doc.add_paragraph()
    return t

def md_to_docx(doc, md: str, box: tuple[str, str] | None = None, heading_base=3):
    """Render section Markdown into the document. box=(fill,border) shades body paragraphs."""
    lines = md.split("\n"); i = 0
    fill, border = box if box else (None, None)
    while i < len(lines):
        line = lines[i].rstrip(); i += 1
        if not line.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            lvl = min(4, heading_base + len(m.group(1)) - 1); heading(doc, render.EMPH_STRIP(m.group(2)) if hasattr(render, "EMPH_STRIP") else m.group(2).replace("**", ""), lvl); continue
        if line.lstrip().startswith("|"):
            tl = [line]
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                tl.append(lines[i]); i += 1
            rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in tl if not set(l.replace("|", "").strip()) <= set("-: ")]
            if rows: add_table(doc, rows[0], rows[1:], fill="D9DEE5" if not fill else fill)
            continue
        fm = render.FIG.match(line.strip())
        if fm:
            b = render.figure_asset(fm.group(1).strip())
            if b and b.get("asset") and (ROOT / b["asset"]).exists():
                try:
                    doc.add_picture(str(ROOT / b["asset"]), width=Mm(150)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    pass
            if fm.group(2).strip(): para(doc, fm.group(2).strip(), style="KDJ Caption", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if render.ARTIFACT_LINE.match(line):
            continue
        fn = render.FOOT.match(line.strip())
        if fn:
            para(doc, f"{fn.group(1) or '•'}. {fn.group(2)}", size=9, space_after=2); continue
        lm = re.match(r"^\s*-\s+(.*)$", line)
        if lm:
            p = para(doc, lm.group(1), style="List Bullet", fill=fill, border=border); continue
        om = re.match(r"^\s*\d+[\.\)]\s+(.*)$", line)
        if om:
            para(doc, om.group(1), style="List Number", fill=fill, border=border); continue
        cm = re.match(r"^\*(.+)\*$", line.strip())
        if cm and len(line) < 320:
            para(doc, cm.group(1), style="KDJ Caption", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER); continue
        # paragraph: join following non-empty non-special lines
        buf = [line.strip()]
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||\s*-\s|\s*\d+[\.\)]\s|\[FIGURE|\[\^)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        para(doc, " ".join(buf), fill=fill, border=border)

def ensure_styles(doc):
    st = doc.styles
    normal = st["Normal"]; normal.font.name = FA_FONT; normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts")) or OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), LATIN_FONT); rf.set(qn("w:hAnsi"), LATIN_FONT); rf.set(qn("w:cs"), FA_FONT); rpr.append(rf)
    normal.paragraph_format.line_spacing = 1.5; normal.paragraph_format.space_after = Pt(6)
    for lvl, sz in ((1, 20), (2, 16), (3, 13.5), (4, 12)):
        h = st[f"Heading {lvl}"]; h.font.name = FA_FONT; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0x1B, 0x1F, 0x24)
        h.paragraph_format.space_before = Pt(18 if lvl <= 2 else 12); h.paragraph_format.keep_with_next = True
    from docx.enum.style import WD_STYLE_TYPE
    for name in ["KDJ Caption", "KDJ Code"] + [v[0] for v in BOX_STYLES.values()]:
        if name not in [s.name for s in st]:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH); s.base_style = st["Normal"]; s.font.size = Pt(10.5 if name != "KDJ Code" else 9)
            if name == "KDJ Code": s.font.name = MONO_FONT
            if name == "KDJ Caption": s.font.italic = True; s.font.size = Pt(9.5)

def add_custom_properties(doc, props: dict):
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"; vt = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    items = "".join(f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i+2}" name="{k}"><vt:lpwstr>{str(v).replace("&","&amp;").replace("<","&lt;")}</vt:lpwstr></property>' for i, (k, v) in enumerate(props.items()))
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="{ns}" xmlns:vt="{vt}">{items}</Properties>'
    part = Part(PackURI("/docProps/custom.xml"), "application/vnd.openxmlformats-officedocument.custom-properties+xml", xml.encode("utf-8"), doc.part.package)
    doc.part.package.relate_to(part, RT.CUSTOM_PROPERTIES)

def add_toc(doc):
    p = doc.add_paragraph(); set_rtl(p)
    r = p.add_run(); fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin"); r._r.append(fld)
    r2 = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = 'TOC \\o "1-2" \\h \\z \\u'; r2._r.append(it)
    r3 = p.add_run(); fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate"); r3._r.append(fld2)
    r4 = p.add_run("فهرست مطالب پس از باز کردن فایل در Word و به‌روزرسانی فیلد (F9) نمایش داده می‌شود."); set_run_fonts(r4, size=10)
    r5 = p.add_run(); fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end"); r5._r.append(fld3)

def build():
    book = json.load(open(ROOT / "master" / "book.json", encoding="utf-8"))
    ver = json.load(open(ROOT / "provenance" / "verification.json", encoding="utf-8")) if (ROOT / "provenance" / "verification.json").exists() else {}
    doc = Document(); ensure_styles(doc)
    sec = doc.sections[0]; sec.page_width, sec.page_height = Mm(210), Mm(297); sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = sec.bottom_margin = Mm(25)
    sectPr = sec._sectPr; bidi = OxmlElement("w:bidi"); sectPr.append(bidi)
    # title page
    cover = ROOT / "images" / "covers" / "IMG-COVER.png"
    if cover.exists():
        doc.add_picture(str(cover), width=Mm(110)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        for _ in range(6): doc.add_paragraph()
    para(doc, book["title_fa"], size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    para(doc, book["title_en"], size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4): doc.add_paragraph()
    para(doc, "ترجمه، تدوین و بازآرایی فارسی: " + PUB["compiler"], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "بر پایه اسناد و انتشارات اصلی NIST و منابع مشخص‌شده در هر فصل", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "وب‌سایت: kohandezh.com", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"ویرایش {PUB['edition_id']}", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()
    # publication page
    heading(doc, "شناسنامهٔ انتشار", 1)
    rows = [("عنوان", book["title_fa"]), ("مترجم / تدوینگر", PUB["compiler"]), ("ناشر", "کهن‌دژ — kohandezh.com"), ("ویرایش", PUB["edition_id"]), ("شناسهٔ انتشار", ver.get("release_id", "Release Candidate")),
            ("تاریخ ساخت", book["built_at"]), ("ریشهٔ مرکل", book["merkle_root"]), ("مشخصات هش", PROV["spec_version"]), ("وضعیت امضا", ver.get("signing_status", "PUBLISHER_SIGNING_KEY_REQUIRED"))]
    add_table(doc, ["فیلد", "مقدار"], [[k, v] for k, v in rows])
    para(doc, "این اثر ترجمهٔ رسمی NIST نیست و NIST آن را تأیید یا منتشر نکرده است. متن اصلی هر سند به نویسندگان و ناشر اصلی تعلق دارد. بخش‌های تدوینگر و لایهٔ اجرایی ایرانی با برچسب مشخص از متن ترجمه‌شده متمایز شده‌اند.", size=10)
    doc.add_page_break()
    heading(doc, "فهرست مطالب", 1); add_toc(doc); doc.add_page_break()
    # preface / how to use (book-level editorial, if generated)
    for kind, title in (("preface", "پیشگفتار"), ("how_to_use", "راهنمای استفاده از کتاب")):
        p = ROOT / "work" / "editorial" / "BOOK" / f"{kind}.json"
        if p.exists():
            ed = json.load(open(p, encoding="utf-8")); heading(doc, ed.get("title_fa") or title, 1)
            for s in ed["sections"]:
                heading(doc, s["title_fa"], 2); md_to_docx(doc, s["fa_text"], box=("FFF8E6", "E0A800"))
            doc.add_page_break()
    # body
    for P in book["parts"]:
        for _ in range(8): doc.add_paragraph()
        para(doc, P["part_id"], size=16, align=WD_ALIGN_PARAGRAPH.CENTER); heading(doc, P["title_fa"], 0 if False else 1)
        para(doc, "منابع اصلی: " + "، ".join(SRC[d].get("series_identifier", d) for d in P["sources"]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        doc.add_page_break()
        for C in P["chapters"]:
            if not C["sections"]: continue
            heading(doc, C["title_fa"], 1)
            if C["kind"] != "source":
                para(doc, f"[{render.LABELS.get(C['origin'], 'یادداشت تدوینگر')}] این فصل محتوای تدوینگر است و متن NIST نیست.", size=9.5, fill="FFF8E6", border="E0A800")
            else:
                d = SRC.get(C["source_doc"], {}); pages = sorted({p for S in C["sections"] for p in S["source_pages"]})
                add_table(doc, ["منبع اصلی", ""], [["شناسه سند", d.get("series_identifier", "")], ["عنوان اصلی", d.get("title", "")], ["نویسندگان", "، ".join(d.get("authors") or []) or d.get("organization", "")],
                                                    ["نوع / وضعیت انتشار", f"{d.get('publication_type','')} — {d.get('publication_status','')}"], ["تاریخ", d.get("publication_date", "")], ["صفحات استفاده‌شده", f"{pages[0]}–{pages[-1]}" if pages else "—"], ["DOI", d.get("doi", "—")]])
            for S in C["sections"]:
                heading(doc, S["title_fa"], 2, cid=S["content_id"])
                if S["origin"] == "source_translation":
                    md_to_docx(doc, S["fa_text"])
                else:
                    lbl = S.get("label_fa") or render.LABELS.get(S["origin"], "یادداشت تدوینگر")
                    key = "executive" if lbl == "برای مدیران" else "technical" if lbl == "برای تیم فنی" else "checklist" if "چک‌لیست" in lbl else "scenario" if "سناریو" in lbl else "template" if "فرم" in lbl else S["origin"]
                    style, fill, border = BOX_STYLES.get(key, BOX_STYLES["editorial_synthesis"])
                    para(doc, lbl, style=style, size=10, bold=True, fill=fill, border=border, space_after=2)
                    md_to_docx(doc, S["fa_text"], box=(fill, border))
                    para(doc, "محتوای تدوینگر — متن NIST نیست. مبتنی بر: " + "، ".join(S.get("derived_from_content_ids", [])[:5]), size=8.5, fill=fill, border=border)
            doc.add_page_break()
    # appendices
    heading(doc, "پیوست الف — واژه‌نامه", 1)
    gp = ROOT / "translation_memory" / "glossary.json"
    if gp.exists():
        terms = sorted([t for t in json.load(open(gp, encoding="utf-8"))["terms"].values() if t.get("preferred_fa")], key=lambda t: t["preferred_fa"])
        add_table(doc, ["فارسی", "English", "تعریف", "وضعیت"], [[t["preferred_fa"], t["english"] + (f" ({t['acronym']})" if t.get("acronym") else ""), t.get("definition_fa") or "", t["status"]] for t in terms])
        heading(doc, "پیوست ب — اختصارات", 1)
        add_table(doc, ["اختصار", "English", "فارسی"], [[t["acronym"], t["english"], t["preferred_fa"]] for t in terms if t.get("acronym")])
    heading(doc, "پیوست ج — کتاب‌شناسی منابع", 1)
    for d in SRC.values():
        para(doc, f"{d.get('series_identifier','')} — {d.get('title','')}. {', '.join(d.get('authors') or []) or d.get('organization','')}. {d.get('publication_type','')}, {d.get('publication_status','')}, {d.get('publication_date','')}. DOI: {d.get('doi') or '—'}. SHA-256: {d['sha256'][:16]}…", size=10)
    heading(doc, "پیوست د — نقشهٔ منابع", 1)
    smap = json.load(open(ROOT / "master" / "source_map.json", encoding="utf-8")) if (ROOT / "master" / "source_map.json").exists() else []
    add_table(doc, ["شناسهٔ بخش", "سند", "بخش منبع", "صفحات", "لایه"], [[r["structural_id"], r["source_series"], r["source_section"] or "", (f"{r['source_pages'][0]}–{r['source_pages'][-1]}" if r["source_pages"] else "—"), render.LABELS.get(r["origin"], r["origin"])] for r in smap])
    heading(doc, "پیوست ه — اصالت و اعتبارسنجی", 1)
    para(doc, f"هر بخش این کتاب یک شناسهٔ محتوا (Content ID) و یک هش SHA-256 دارد که در درخت مرکل ویرایش {PUB['edition_id']} ثبت شده است. ریشهٔ مرکل: ", size=10)
    para(doc, book["merkle_root"], mono=True)
    para(doc, "برای بررسی اصالت، به kohandezh.com/ai-book/verify/ مراجعه کنید.", size=10)
    add_custom_properties(doc, {"KDJ Edition ID": PUB["edition_id"], "Release ID": ver.get("release_id", "RC"), "Master Hash": book.get("merkle_root", ""), "Publisher": "کهن‌دژ", "Translator/Compiler": PUB["compiler"],
                                "Website": PUB["website"], "Build Date": book["built_at"], "Manifest ID": ver.get("manifest_sha256", "pending"), "Hash Spec": PROV["spec_version"]})
    cp = doc.core_properties; cp.title = book["title_fa"]; cp.author = PUB["compiler"]; cp.subject = "AI governance, security and risk management — Persian edition based on NIST publications"; cp.keywords = "NIST; AI RMF; هوش مصنوعی; کهن‌دژ"; cp.language = "fa-IR"; cp.category = PUB["edition_id"]
    out = ROOT / "docx" / "fa" / "book-fa.docx"; out.parent.mkdir(parents=True, exist_ok=True); doc.save(out)
    Document(out)  # re-open check
    db.event("docx_built", {"path": str(out.relative_to(ROOT)), "bytes": out.stat().st_size})
    print("DOCX built:", out, out.stat().st_size, "bytes")

if __name__ == "__main__":
    build()
